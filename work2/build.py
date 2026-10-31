#!/usr/bin/env python3
"""Build the Legal Reach (85 listings) deliverables: Excel + CSVs."""
import csv, datetime, os, re
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.abspath(os.path.join(HERE, ".."))
SOURCE = ("https://www.legalreach.com/search_results?tid=48&location_value=United%20States"
          "&place_id=ChIJCzYy5IS16lQRQrfeQ5K5Oxw")

# kind: F=firm, I=individual (person-centric listing naming a firm), D=duplicate/office listing,
#       X=foreign / non-US, U=unclear
# (n, listing name as published, city, state, url, kind, canonical firm, website, evidence)
ROWS = [
 (1, "Ryan Clement | Business & Technology Legal Group", "Denver, CO", "https://www.legalreach.com/united-states/denver/attorney/ryan-clement", "I",
  "Business & Technology Legal Group", "https://businessandtechlawyers.com/",
  "Listing is person-led; firm site found by name+city search and matches Denver address on the Legal Reach profile"),
 (2, "Charles L. Riddle | EsquireTrademarks", "Philadelphia, PA", "https://www.legalreach.com/pennsylvania/philadelphia/attorney/charles-l-riddle", "I",
  "EsquireTrademarks", "https://esquiretrademarks.com/",
  "Website published on the firm's own profiles; cross-checked with the firm's LegalServicesLink profile (independently extracted earlier)"),
 (3, "TBillick Law PLLC", "Seattle, WA", "https://www.legalreach.com/washington/seattle/attorney/tbillick-law-pllc", "F",
  "TBillick Law PLLC", "https://www.tbillicklaw.com/", "Official site confirmed by name+city search (multiple directories list the same domain)"),
 (4, "Michael Ahmadshahi, PhD, Law Offices", "Newport Beach, CA", "https://www.legalreach.com/california/newport-beach/attorney/michael-ahmadshahi-phd-law-offices", "F",
  "Michael Ahmadshahi, PhD, Law Offices", "https://mmaiplaw.com/", "Firm site found by search; address on site (4695 MacArthur Ct, Newport Beach) matches the listing"),
 (5, "Bukh Law Firm, PLLC", "Brooklyn, NY", "https://www.legalreach.com/new-york/brooklyn/attorney/bukh-law-firm-pllc", "F",
  "Bukh Law Firm, PLLC", "https://bukhlaw.com/", "Domain confirmed by several independent directories listing the same Brooklyn address/phone"),
 (6, "Corey Pollard Law", "Richmond, VA", "https://www.legalreach.com/virginia/richmond/attorney/corey-pollard-law", "F",
  "Corey Pollard Law", "https://cpollardlaw.com/", "Firm site; domain appears on the firm's own professional listings with the same Richmond phone"),
 (7, "Law Offices of Lloyd S. Mann", "Sherman Oaks, CA", "https://www.legalreach.com/california/sherman-oaks/attorney/law-offices-of-lloyd-s-mann", "F",
  "Law Offices of Lloyd S. Mann", "https://www.themannlaw.com/", "Website listed for the same Sherman Oaks office address on Martindale/FindLaw"),
 (8, "Scalfone Law PLLC", "Syracuse, NY", "https://www.legalreach.com/syracuse/attorney/scalfone-law-pllc", "F",
  "Scalfone Law PLLC", "https://www.scalfonelaw.com/", "Website published on the Legal Reach profile itself"),
 (9, "The Kinder Law Group", "Irvine, CA", "https://www.legalreach.com/irvine/attorney/the-kinder-law-group", "F",
  "The Kinder Law Group", "https://tklglaw.com/", "Firm site; matches Irvine address and the listing's IP boutique description"),
 (10, "Adam S Tracy | The Tracy Firm, Ltd.", "Chicago, IL", "https://www.legalreach.com/chicago/attorney/adam-s-tracy", "I",
  "The Tracy Firm, Ltd.", "https://tracyfirm.com/",
  "Cross-checked with the firm's LegalServicesLink profile (independently extracted earlier)"),
 (11, "Adler Pollock & Sheehan P.C.", "Providence, RI", "https://www.legalreach.com/rhode-island/providence/attorney/adler-pollock-sheehan-pc", "F",
  "Adler Pollock & Sheehan P.C.", "https://www.apslaw.com/", "Firm site; Rhode Island HQ and 401-274-7200 match"),
 (12, "Adler Pollock & Sheehan, PC", "Providence, RI", "https://www.legalreach.com/rhode-island/providence/attorney/adler-pollock-sheehan-pc-634799", "D",
  "Adler Pollock & Sheehan P.C.", "", "Second profile for the same Providence firm (same address/phone) - merged into #11"),
 (13, "Ahlawat & Associates Advocates", "Irvine, CA (India firm)", "https://www.legalreach.com/united-states/irvine/attorney/ahlawat-associates-advocates", "X",
  "Ahlawat & Associates", "https://www.ahlawatassociates.com/", "Non-US: profile describes an Indian law firm; New Delhi HQ"),
 (14, "Alistair Chan IP", "Bainbridge Island, WA", "https://www.legalreach.com/united-states/bainbridge-island/attorney/alistair-chan-ip", "F",
  "Alistair Chan IP", "https://achanip.legal/", "Website published in the Legal Reach listing description"),
 (15, "Allen, Dyer, Doppelt + Gilchrist, PA", "Miami, FL", "https://www.legalreach.com/united-states/miami/attorney/allen-dyer-doppelt-gilchrist-pa", "F",
  "Allen, Dyer, Doppelt + Gilchrist, P.A.", "https://allendyer.com/", "Firm site; Miami Brickell Ave office matches"),
 (16, "Bailey & Galyen Attorneys at Law", "Bedford, TX", "https://www.legalreach.com/united-states/bedford/attorney/bailey-galyen-attorneys-at-law", "F",
  "Bailey & Galyen Attorneys at Law", "https://www.thetexasattorney.com/", "Firm site; Bedford TX HQ address matches"),
 (17, "Bold IP, PLLC", "Seattle, WA", "https://www.legalreach.com/washington/seattle/attorney/bold-ip-pllc", "F",
  "Bold IP, PLLC", "https://boldip.com/", "Firm site (also trades as Bold Patents); Washington-based IP firm"),
 (18, "Burdine Law", "Nashville, TN", "https://www.legalreach.com/united-states/nashville/attorney/burdine-law-645871", "F",
  "Burdine Law", "https://www.burdinelawfirm.com/", "Website published in the Legal Reach listing description"),
 (19, "Chase Law & Associates, P.A.", "Miami Beach, FL", "https://www.legalreach.com/united-states/miami-beach/attorney/chase-law-associates-p-a", "F",
  "Chase Law & Associates, P.A.", "https://chaselaw.com/", "Cross-checked with the firm's LegalServicesLink profile (independently extracted earlier)"),
 (20, "Christensen, Fonder, Dardi & Herbert PLLC", "Minneapolis, MN", "https://www.legalreach.com/united-states/minneapolis/attorney/christensen-fonder-dardi-herbert-pllc", "F",
  "Christensen, Fonder, Dardi & Herbert PLLC", "https://www.cfd-ip.com/", "Cross-checked with the firm's LegalServicesLink profile (independently extracted earlier)"),
 (21, "Clayton Trial Lawyers LP", "Fort Lauderdale, FL", "https://www.legalreach.com/united-states/fort-lauderdale/attorney/clayton-trial-lawyers-lp", "F",
  "Clayton Trial Lawyers", "https://ctllawyers.com/", "Firm site; 401 E Las Olas Blvd address and 833-938-2300 phone match"),
 (22, "Collard & Roe, P.C.", "Roslyn, NY", "https://www.legalreach.com/united-states/roslyn/attorney/collard-roe-p-c", "F",
  "Collard & Roe, P.C.", "https://collardroe.com/", "Firm site; 1077 Northern Blvd Roslyn address matches"),
 (23, "Creek & River, Inc.", "Irvine, CA", "https://www.legalreach.com/irvine/attorney/creek-20-river-inc", "F",
  "Creek & River Global, Inc.", "https://www.creekandriverglobal.com/", "US arm of Japan's Creek & River Co. (parent cri.co.jp), Irvine office. Note: a legal-service consulting / attorney-introduction business rather than a traditional law firm"),
 (24, "Curley & Rothman, LLC", "Conshohocken, PA", "https://www.legalreach.com/pennsylvania/conshohocken/attorney/curley-20-rothman-llc", "F",
  "Curley & Rothman, LLC", "https://www.curleyrothman.com/", "Firm site; Spring Mill Corporate Center address matches"),
 (25, "Dennen Ann | Dennen IP Law", "Huntsville, AL", "https://www.legalreach.com/united-states/huntsville/attorney/dennen-ann", "I",
  "Dennen IP Law, LLC", "https://www.denneniplaw.com/", "Firm site; Huntsville IP practice, same phone as the listing"),
 (26, "Diament Patent Law, P.C.", "Encino, CA", "https://www.legalreach.com/united-states/encino/attorney/diament-patent-law-p-c", "F",
  "Diament Patent Law, P.C.", "https://www.diamentpatentlaw.com/", "Cross-checked with the firm's LegalServicesLink profile (independently extracted earlier)"),
 (27, "DiAngelo Law", "Tampa, FL", "https://www.legalreach.com/florida/tampa/attorney/diangelo-law", "F",
  "DiAngelo Law", "https://diangelolaw.com/", "Firm site; 1112 Channelside Dr Tampa address matches"),
 (28, "Dutch Law Firm in the Netherlands", "Amsterdam (NL)", "https://www.legalreach.com/new-york/amsterdam/attorney/dutch-law-firm-in-the-netherlands", "X",
  "MAAK Advocaten NV (Netherlands)", "", "Non-US: Dutch firm; the directory filed it under a New York 'Amsterdam' locality"),
 (29, "Elliott Alderman | The Alderman Law Office", "Washington, DC", "https://www.legalreach.com/united-states/washington/attorney/elliott-alderman", "I",
  "The Alderman Law Office", "https://thecontentlawyer.com/", "Firm site; 1455 Pennsylvania Ave NW and 202-596-5682 match the listing"),
 (30, "Ford Firm", "Tuscaloosa, AL", "https://www.legalreach.com/united-states/tuscaloosa/attorney/ford-firm", "F",
  "Ford Firm", "https://fordfirmlaw.com/", "Firm site; 705 27th Ave Tuscaloosa address matches"),
 (31, "Foreclosure HELP", "New York, NY", "https://www.legalreach.com/new-york/new-yourk-city/attorney/foreclosure-help", "U",
  "Foreclosure HELP", "", "Unclear: real-estate/employment lead-generation listing with a non-IP promotional URL; no verifiable official firm site"),
 (32, "Geoffrey Gussis | Geoffrey G. Gussis, Esq.", "New Brunswick, NJ", "https://www.legalreach.com/united-states/new-brunswick/attorney/geoffrey-gussis", "I",
  "Geoffrey G. Gussis, LLC", "https://www.gussislaw.com/", "Firm site; 83 Morris St New Brunswick address matches"),
 (33, "Geonetta & Frucht, LLP", "San Francisco, CA", "https://www.legalreach.com/united-states/san-francisco/attorney/geonetta-frucht-llp", "F",
  "Geonetta & Frucht, LLP", "https://www.geonetta-frucht.com/", "Firm site; 100 Montgomery St address matches"),
 (34, "Hadi Harp", "Los Angeles, CA", "https://www.legalreach.com/los-angeles/attorney/hadi-harp", "I",
  "Harp Law", "https://lawharp.com/", "Solo practice listed as Harp Law; same Los Angeles phone as the listing"),
 (35, "Halloran Sage", "Hartford, CT", "https://www.legalreach.com/hartford/attorney/halloran-sage", "F",
  "Halloran Sage", "https://www.halloran-sage.com/", "Firm site; One Goodwin Square Hartford and 860-522-6103 match"),
 (36, "HamptonHolley LLP", "Corona Del Mar, CA", "https://www.legalreach.com/california/corona-del-mar/attorney/hamptonholley-llp", "F",
  "HamptonHolley LLP", "https://www.hamptonholley.com/", "Firm site; 2101 E Coast Hwy Corona del Mar address matches"),
 (37, "Herman Legal Group, LLC", "McKinney, TX", "https://www.legalreach.com/united-states/mckinney/attorney/herman-legal-group-llc", "F",
  "Herman Legal Group, LLC", "https://www.lawfirm4immigrants.com/", "Firm site; the firm's multi-office listings share the same brand/phone pattern"),
 (38, "Herman Legal Group, LLC", "Charlotte, NC", "https://www.legalreach.com/north-carolina/charlotte/attorney/herman-legal-group-llc", "D",
  "Herman Legal Group, LLC", "", "Additional office listing of the same firm (Charlotte) - merged into #37"),
 (39, "Herman Legal Group, LLC", "Painesville, OH", "https://www.legalreach.com/united-states/painesville/attorney/herman-legal-group-llc", "D",
  "Herman Legal Group, LLC", "", "Additional office listing of the same firm (Painesville OH) - merged into #37"),
 (40, "HPS Law Group LLC", "Mesa, AZ", "https://www.legalreach.com/united-states/mesa/attorney/hps-law-group-llc", "F",
  "HPS Law Group LLC", "https://hpslawgroup.com/", "Firm site; 123 N Centennial Way Mesa address matches"),
 (41, "IndiaLaw LLP", "Mumbai (IN)", "https://www.legalreach.com/mumbai/attorney/indialaw-llp", "X",
  "IndiaLaw LLP", "", "Non-US: Indian firm (Mumbai, +91 22 number); no US office"),
 (42, "Infinity Patent Services", "Riverside, CA", "https://www.legalreach.com/united-states/riverside/attorney/infinity-patent-services-639546", "F",
  "Infinity Patent Services", "https://www.saenzlawoffice.com/", "Cross-checked with the firm's LegalServicesLink profile (independently extracted earlier)"),
 (43, "IP RANI - The Law Firm", "Surajpur, Uttar Pradesh (IN)", "https://www.legalreach.com/india/surajpur/attorney/ip-rani-the-law-firm", "X",
  "IP RANI - The Law Firm", "", "Non-US: Delhi NCR (India) firm, listed in the US directory by mistake"),
 (44, "Israel & Gerity, PLLC", "Phoenix, AZ", "https://www.legalreach.com/united-states/phoenix/attorney/israel-gerity-pllc", "F",
  "Israel & Gerity, PLLC", "https://www.ig-law.com/", "Firm site (main site listed on the firm's FindLaw/Martindale profiles); Phoenix 85012 address matches"),
 (45, "Jahn & Associates", "Boulder, CO", "https://www.legalreach.com/colorado/boulder/attorney/jahn-associates", "F",
  "Jahn & Associates", "https://www.jahnlaw.com/", "Firm site; 1942 Broadway Boulder address matches"),
 (46, "John C. Laurence Law, PLLC", "New York, NY", "https://www.legalreach.com/united-states/new-york/attorney/john-c-laurence-law-pllc", "F",
  "John C. Laurence Law, PLLC", "https://www.jclaurencelaw.com/", "Firm site; same 325 E 79th St office (the practice also uses TCP Law, PLLC / tcplawfirm.com)"),
 (47, "Klemchuk LLP", "Dallas, TX", "https://www.legalreach.com/texas/dallas/attorney/klemchuk-llp", "F",
  "Klemchuk LLP", "https://www.klemchuk.com/", "Firm site; 8150 N Central Expy Dallas and 214-367-6000 match"),
 (48, "Law Office of Dorian Cartwright", "San Jose, CA", "https://www.legalreach.com/san-jose/attorney/law-office-of-dorian-cartwright-644289", "F",
  "Law Office of Dorian Cartwright", "https://www.cartwrightesq.com/", "Firm site; 2130 The Alameda San Jose address matches"),
 (49, "Law Office of Dorian Cartwright", "San Jose, CA", "https://www.legalreach.com/san-jose/attorney/law-office-of-dorian-cartwright", "D",
  "Law Office of Dorian Cartwright", "", "Second profile for the same San Jose solo firm - merged into #48"),
 (50, "Law Office of Robert L. Greener, P.C.", "New York, NY", "https://www.legalreach.com/united-states/new-york/attorney/law-office-of-robert-l-greener-p-c", "F",
  "Law Office of Robert L. Greener, P.C.", "https://greenerlegal.com/", "Website published in the Legal Reach listing description"),
 (51, "Law Offices of David P. Gardner", "Santa Rosa, CA", "https://www.legalreach.com/united-states/santa-rosa/attorney/law-offices-of-david-p-gardner", "F",
  "Law Offices of David P. Gardner", "https://davidpgardnerlaw.com/", "Firm site; same Santa Rosa/Petaluma practice and 707-658-4460 phone"),
 (52, "Law Offices of Edward C. Chen", "Irvine, CA", "https://www.legalreach.com/united-states/irvine/attorney/law-offices-of-edward-c-chen-639924", "F",
  "Law Offices of Edward C. Chen", "https://edchenlaw.com/", "Firm site; 1 Park Plaza Irvine address matches"),
 (53, "Law Offices of Paul Share", "New York, NY", "https://www.legalreach.com/united-states/new-york/attorney/law-offices-of-paul-share", "F",
  "Law Offices of Paul Share", "https://www.sharelawfirm.com/", "Firm site; 236 W 30th St New York and 212-643-0700 match"),
 (54, "Lawyers in America", "New York, NY", "https://www.legalreach.com/new-york/new-york-city/attorney/lawyers-in-america", "U",
  "Lawyers in America", "", "Unclear: Russian-language attorney-referral/lead service, not an identifiable law firm; no verifiable official site"),
 (55, "Libenzon & Associates", "Boston, MA", "https://www.legalreach.com/united-states/boston/attorney/libenzon-associates", "F",
  "Libenzon & Associates", "https://www.timetopatent.com/", "Cross-checked with the firm's LegalServicesLink profile (independently extracted earlier)"),
 (56, "Lloyd & Mousilli", "Houston, TX", "https://www.legalreach.com/united-states/houston/attorney/lloyd-mousilli", "F",
  "Lloyd & Mousilli", "https://www.lloydmousilli.com/", "Firm site; 11807 Westheimer Rd Houston address matches"),
 (57, "Long and Associates", "Los Angeles, CA", "https://www.legalreach.com/california/los-angeles/attorney/long-and-associates", "F",
  "Long and Associates", "https://aexius.com/", "The firm's own profiles point to aexius.com and its Los Angeles street address (1920 Hillhurst Ave) matches the listing; the domain does not match the firm name, so this is the one entry worth a manual double-check"),
 (58, "Marc Randazza | Randazza Legal Group", "Las Vegas, NV", "https://www.legalreach.com/united-states/las-vegas/attorney/marc-randazza", "I",
  "Randazza Legal Group", "https://randazza.com/", "Firm site; Las Vegas HQ matches"),
 (59, "Mark Shield", "Delhi (IN)", "https://www.legalreach.com/delhi/attorney/mark-shield", "X",
  "Mark Shield", "", "Non-US: Indian IP firm (Delhi), listed in the US directory by mistake"),
 (60, "Milwaukee Personal Injury Lawyer Group", "Milwaukee, WI", "https://www.legalreach.com/wisconsin/milwaukee/attorney/milwaukee-personal-injury-lawyer-group", "F",
  "Milwaukee Personal Injury Lawyer Group", "https://www.milwaukeeinjurylawgroup.com/", "Site listed for the same 310 E Buffalo St Milwaukee office and phone"),
 (61, "NY Business and Corporate Law - Ishimbayev Law Firm", "New York, NY", "https://www.legalreach.com/new-york/new-york/attorney/ny-business-and-corporate-law-ishimbayev-law-firm", "F",
  "Ishimbayev Law Firm, P.C.", "https://www.ishimbayev.com/", "Firm site; 1 World Trade Center Suite 8500 address matches"),
 (62, "Omni Legal Group", "Los Angeles, CA", "https://www.legalreach.com/california/los-angeles/attorney/omni-legal-group", "F",
  "Omni Legal Group", "https://omnilegalgroup.com/", "Cross-checked with the firm's LegalServicesLink profile (independently extracted earlier)"),
 (63, "Paley Legal", "Los Angeles, CA", "https://www.legalreach.com/united-states/los-angeles/attorney/paley-legal", "F",
  "Paley Legal", "https://www.paleylegal.com/", "Firm site; Los Angeles practice of Jennifer Paley, same phone as the listing"),
 (64, "Perkins Pershes, PLLC", "Boca Raton, FL", "https://www.legalreach.com/united-states/boca-raton/attorney/perkins-pershes-pllc", "F",
  "Perkins Pershes, PLLC", "https://www.perkinspershes.com/", "Firm site (brand-protection practice also uses brandprotection.law)"),
 (65, "Randazza Legal Group Gloucester Office", "Gloucester, MA", "https://www.legalreach.com/united-states/gloucester/attorney/randazza-legal-group-gloucester-office", "D",
  "Randazza Legal Group", "", "Branch-office listing of the same firm (Gloucester MA) - merged into #58"),
 (66, "Reilly Intellectual Property Law", "Denver, CO", "https://www.legalreach.com/united-states/denver/attorney/reilly-intellectual-property-law", "F",
  "Reilly Intellectual Property Law Firm, P.C.", "https://www.patentattorneysdenver.com/", "Firm site; 303-839-8700 phone matches the listing"),
 (67, "Ryan Smith & Carbine, Ltd.", "Rutland, VT", "https://www.legalreach.com/united-states/rutland/attorney/ryan-smith-carbine-ltd", "F",
  "Ryan Smith & Carbine, Ltd.", "https://www.rsclaw.com/", "Cross-checked with the firm's LegalServicesLink profile (independently extracted earlier)"),
 (68, "Seiter Legal Studio", "Santa Monica, CA", "https://www.legalreach.com/united-states/santa-monica/attorney/seiter-legal-studio", "F",
  "Seiter Legal Studio", "https://www.seiterlegalstudio.com/", "Firm site; 2450 Colorado Ave Santa Monica address matches"),
 (69, "Shumaker, Loop & Kendrick, LLP", "Toledo, OH", "https://www.legalreach.com/united-states/toledo/attorney/shumaker-loop-kendrick-llp", "F",
  "Shumaker, Loop & Kendrick, LLP", "https://www.shumaker.com/", "Firm site; 1000 Jackson St Toledo and 419-241-9000 match"),
 (70, "Stanzione & Associates, PLLC", "New York, NY", "https://www.legalreach.com/united-states/new-york/attorney/stanzione-associates-pllc", "F",
  "Stanzione & Associates, PLLC", "https://www.stanzioneiplaw.com/", "Firm site; 477 Madison Ave New York address matches"),
 (71, "Stephen L. Keefe LLC", "Mahwah, NJ", "https://www.legalreach.com/united-states/mahwah/attorney/stephen-l-keefe-llc", "F",
  "Stephen L. Keefe LLC (Keefe IP)", "https://www.keefeip.com/", "Website published in the Legal Reach listing description"),
 (72, "The Berggren Law Office", "Oakdale, MN", "https://www.legalreach.com/united-states/oakdale/attorney/the-berggren-law-office", "F",
  "Berggren Law Office, LLC", "https://berggrenlaw.com/", "Firm site; 7090 43rd St N Oakdale and 651-748-5998 match (alternative domain wrberggren.com)"),
 (73, "The Corcoran Law Firm Pllc", "Texarkana, TX", "https://www.legalreach.com/united-states/texarkana/attorney/the-corcoran-law-firm-pllc", "F",
  "The Corcoran Law Firm, PLLC", "https://www.corcoranip.com/", "Firm site; IP practice site named on the listing profile itself"),
 (74, "The Farber Law Firm", "Coral Gables, FL", "https://www.legalreach.com/coral-gables/attorney/the-farber-law-firm", "F",
  "The Farber Law Firm", "https://www.thefarberlawfirm.com/", "Firm site; 2199 Ponce de Leon Blvd and 305-774-0134 match"),
 (75, "The Law Offices of Bryan J. Reddix", "Washington, DC", "https://www.legalreach.com/district-of-columbia/washington/attorney/the-law-offices-of-bryan-j-reddix", "F",
  "The Law Offices of Bryan J. Reddix", "", "No official firm website found; left blank rather than guessed"),
 (76, "The Law Offices of Doroshow, Pasquale, Krawitz & Bhaya", "Smyrna, DE", "https://www.legalreach.com/delaware/smyrna/attorney/the-law-offices-of-doroshow-pasquale-krawitz-bhaya", "F",
  "The Law Offices of Doroshow, Pasquale, Krawitz & Bhaya", "https://www.dplaw.com/", "Firm site; Delaware offices and 302-508-2140 match"),
 (77, "The Law Offices of Zakari A. Kurtz", "Sterling, VA", "https://www.legalreach.com/united-states/sterling/attorney/the-law-offices-of-zakari-a-kurtz", "F",
  "The Law Offices of Zakari A. Kurtz", "https://zakurtzlaw.com/", "Firm site named on the firm's own profile page"),
 (78, "The Myers Law Group", "Newport Beach, CA", "https://www.legalreach.com/united-states/newport-beach/attorney/the-myers-law-group", "F",
  "The Myers Law Group", "https://www.themyerslg.com/", "Firm site; 4695 MacArthur Ct Newport Beach and 949-825-5590 match"),
 (79, "The Myers Law Group", "Newport Beach, CA", "https://www.legalreach.com/newport-beach/attorney/the-myers-law-group", "D",
  "The Myers Law Group", "", "Second profile for the same Newport Beach firm - merged into #78"),
 (80, "TMpros", "Austin, TX", "https://www.legalreach.com/texas/austin/attorney/tmpros", "F",
  "TMpros", "https://tmpros.co/", "Firm site; 106 E 6th St Austin and the listing's phone match"),
 (81, "Tom James Law Office", "Cokato, MN", "https://www.legalreach.com/united-states/cokato/attorney/tom-james-law-office", "F",
  "Law Office of Tom James", "https://tomjameslaw.com/", "Firm site; Cokato MN trademark/copyright practice"),
 (82, "Verhagen Bennett LLP", "Santa Monica, CA", "https://www.legalreach.com/united-states/santa-monica/attorney/verhagen-bennett-llp-638274", "F",
  "Verhagen Bennett LLP", "https://www.verhagenbennett.com/", "Firm site; 100 Wilshire Blvd Santa Monica and 310-917-1064 match (a Yelp listing marks the office closed - confirm current status)"),
 (83, "Wagner & Associates", "Miami, FL", "https://www.legalreach.com/florida/miami/attorney/wagner-associates-", "F",
  "Wagner & Associates - The Law Firm", "https://www.danielwagnerlaw.com/", "Firm site; 305-919-7788 phone matches"),
 (84, "WILLIAM SCOTT GOLDMAN | GOLDMAN LAW GROUP, PLLC", "Washington, DC", "https://www.legalreach.com/district-of-columbia/washington/attorney/william-scott-goldman", "I",
  "Goldman Law Group, PLLC", "https://www.goldmanlawgroup.com/", "Firm site; 1300 Pennsylvania Ave NW and 202-880-9200 match"),
 (85, "Xu Law Group", "New York, NY", "https://www.legalreach.com/united-states/new-york/attorney/xu-law-group", "F",
  "Xu Law Group PLLC", "https://www.xulawgroup.com/", "Firm site; One World Trade Center Suite 8500 address matches"),
]

KIND_LABEL = {"F": "Law firm", "I": "Individual attorney", "D": "Duplicate / extra office",
              "X": "Non-US (foreign)", "U": "Unclear / not an IP firm"}

HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(bold=True, color="FFFFFF", size=11)


def sheet(wb, title, headers, rows, widths, table_name, link_cols=()):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font, cell.fill = HDR_FONT, HDR_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    for r in rows:
        ws.append(list(r))
        for col in link_cols:
            v = r[col - 1]
            if v and str(v).startswith("http"):
                cell = ws.cell(row=ws.max_row, column=col)
                cell.hyperlink = v
                cell.font = Font(color="0563C1", underline="single")
    last = ws.max_row
    if last > 1:
        t = Table(displayName=table_name, ref=f"A1:{chr(64 + len(headers))}{last}")
        t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True,
                                          showColumnStripes=False, showFirstColumn=False,
                                          showLastColumn=False)
        ws.add_table(t)
        ws.auto_filter.ref = f"A1:{chr(64 + len(headers))}{last}"
    ws.freeze_panes = "A2"
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = w
    return ws


def load_rows():
    """Join the verified research fields onto the untouched directory capture in work2/all.tsv."""
    tsv = [l.rstrip("\n").split("\t") for l in open(os.path.join(HERE, "all.tsv"), encoding="utf-8")]
    assert len(tsv) == len(ROWS) == 85, (len(tsv), len(ROWS))
    out = []
    for (n, _n, _loc, _url, kind, firm, site, ev), (page, name, url, loc) in zip(ROWS, tsv):
        assert n == len(out) + 1
        out.append((n, name, loc, url, kind, firm, site, ev))
    return out


def main():
    data = load_rows()
    firms, dups, other = {}, [], []
    inds = []
    for (n, name, loc, url, kind, firm, site, ev) in data:
        if kind == "F":
            firms.setdefault(firm, {"site": site, "loc": loc, "listings": [], "ev": ev})
            firms[firm]["listings"].append(n)
        elif kind == "I":
            firms.setdefault(firm, {"site": site, "loc": loc, "listings": [], "ev": ev})
            firms[firm]["listings"].append(n)
            inds.append((name.split(" | ")[0].strip(), firm, site, loc, url, ev))
        elif kind == "D":
            primary = [r for r in data if r[5] == firm and r[4] in ("F", "I")][0]
            firms[firm]["listings"].append(n)
            dups.append((n, name, loc, url, firm, primary[0], ev))
        else:
            other.append((n, name, loc, url, KIND_LABEL[kind], firm, site, ev))

    wb = Workbook()
    wb.remove(wb.active)

    raw = [(n, name, loc, KIND_LABEL[k], url) for (n, name, loc, url, k, *_rest) in data]
    sheet(wb, "Raw Legal Reach Listings",
          ["#", "Listing Name (as published)", "Location", "Listing Type", "Legal Reach URL"],
          raw, [6, 58, 26, 24, 78], "RawListings", link_cols=(5,))

    vrows = []
    for i, (firm, d) in enumerate(sorted(firms.items(), key=lambda kv: kv[1]["listings"][0]), 1):
        vrows.append((i, firm, d["site"], d["loc"],
                      len(d["listings"]),
                      "Website found" if d["site"] else "No website found",
                      d["ev"]))
    sheet(wb, "Verified US IP Firms",
          ["#", "Firm Name", "Official Website", "Location", "Listings Merged", "Website Status", "Evidence / Verification Note"],
          vrows, [6, 48, 44, 26, 12, 18, 96], "VerifiedFirms", link_cols=(3,))

    irows = [(i, a, f, s, loc, url) for i, (a, f, s, loc, url, _e) in enumerate(inds, 1)]
    sheet(wb, "Individual IP Attorneys",
          ["#", "Attorney", "Associated Firm", "Firm Website", "Location", "Legal Reach Profile"],
          irows, [6, 30, 40, 42, 26, 78], "IndividualAttorneys", link_cols=(4, 6))

    orows = [(n, name, loc, k, firm, site, ev) for (n, name, loc, url, k, firm, site, ev) in other]
    sheet(wb, "NIP and Unclear",
          ["#", "Listing Name", "Location", "Reason Category", "As-Listed Firm", "Website", "Notes"],
          orows, [6, 46, 26, 24, 40, 34, 88], "NipUnclear", link_cols=(6,))

    names = {r[0]: r[1] for r in data}
    drow = [(n, name, loc, firm, f"#{p} - {names[p]}", ev) for (n, name, loc, url, firm, p, ev) in dups]
    sheet(wb, "Duplicates",
          ["#", "Duplicate Listing", "Location", "Firm", "Kept Listing", "Why it is a duplicate"],
          drow, [6, 46, 26, 40, 34, 92], "Duplicates")

    ws = wb.create_sheet("Research Summary")
    ws.column_dimensions["A"].width = 60
    ws.column_dimensions["B"].width = 74
    ws["A1"] = "Legal Reach - United States Intellectual Property Directory (85 listings)"
    ws["A1"].font = Font(bold=True, size=14)
    n_firm_sites = sum(1 for d in firms.values() if d["site"])
    rows = [
        ("Primary source (U.S. IP directory, 85 results)", SOURCE),
        ("Pagination used", "&page=1 ... &page=9 - server-rendered, 10 listings per page"),
        ("Directory total reported by the site", "85"),
        ("Listings captured", len(data)),
        ("Directory pages processed", "9 of 9 (10 per page; page 9 holds 5)"),
        ("robots.txt", "Only /api/ disallowed for all agents; directory and profile pages allowed"),
        ("Unique U.S. IP law firms identified", len(firms)),
        ("  - with an official website", n_firm_sites),
        ("  - website not found / left blank", len(firms) - n_firm_sites),
        ("Individual attorneys (person-led listings)", len(inds)),
        ("Duplicate / extra-office listings merged", len(dups)),
        ("Non-US listings found inside the U.S. directory", sum(1 for o in other if o[4] == "Non-US (foreign)")),
        ("Unclear / not an identifiable IP firm", sum(1 for o in other if o[4] == "Unclear / not an IP firm")),
        ("Reconciliation of the 85 listings",
         f"85 = {sum(len(d['listings']) for d in firms.values())} U.S. listings covered by the {len(firms)} practices "
         f"in 'Verified US IP Firms' (of which {len(dups)} are second/extra profiles, detailed on the 'Duplicates' tab) "
         f"+ {sum(1 for o in other if o[4] == 'Non-US (foreign)')} non-US + "
         f"{sum(1 for o in other if o[4] == 'Unclear / not an IP firm')} unclear. Nothing dropped, nothing duplicated."),
        ("Date of research", datetime.date.today().isoformat()),
        ("Method", "Every one of the 9 directory pages was read directly. Each listing's website was taken (a) from the "
                   "Legal Reach profile text where the site publishes one, (b) from a cross-check against the firm's "
                   "LegalServicesLink profile already extracted in this project, or (c) from a name + city web search whose "
                   "result matched the listed address/phone. No URL was guessed; where nothing verifiable was found the cell is blank."),
        ("Caveats", "Legal Reach mixes firms, solo attorneys and extra office listings, and it also carries a few non-US "
                    "firms. Two listings (Verhagen Bennett, Long and Associates) carry status notes worth a manual look."),
    ]
    r = 3
    for label, val in rows:
        ws.cell(row=r, column=1, value=label).font = Font(bold=True)
        ws.cell(row=r, column=2, value=val).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=r, column=1).alignment = Alignment(vertical="top")
        r += 1

    xlsx = os.path.join(OUT, "LegalReach_US_IP_Firms_Research.xlsx")
    wb.save(xlsx)

    with open(os.path.join(OUT, "LegalReach_Verified_US_IP_Firms.csv"), "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["Firm Name", "Official Website", "Location", "Listings Merged", "Website Status"])
        for (i, firm, site, loc, cnt, status, ev) in vrows:
            w.writerow([firm, site, loc, cnt, status])

    with open(os.path.join(OUT, "LegalReach_Raw_Listings_85.csv"), "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["#", "Listing Name", "Location", "Listing Type", "Legal Reach URL"])
        w.writerows(raw)

    print(f"listings={len(data)} firms={len(firms)} with_site={n_firm_sites} "
          f"no_site={len(firms)-n_firm_sites} attorneys={len(inds)} dups={len(dups)} other={len(other)}")
    print("wrote", xlsx)


if __name__ == "__main__":
    main()
