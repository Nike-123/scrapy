import json, os, sys
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
sys.path.insert(0, "/tmp/venv/lib/python3.11/site-packages")
import pytest
from cornell_ip.dedup import deduplicate, normalize_firm_name, website_domain

FIX = os.path.join(os.path.dirname(__file__), "fixtures")

def test_normalize():
    assert normalize_firm_name("Morris, Nichols, Arsht & Tunnell LLP") == "morris nichols arsht and tunnell"
    assert normalize_firm_name("Smith & Jones, P.C.") == "smith and jones"
    assert website_domain("http://www.SVLaw.com/attorneys/x/") == "svlaw.com"
    assert website_domain("") == ""

def test_dedup_first_wins_no_similar_merge():
    recs = [
        {"firm_name": "Alpha Law LLP", "website": "https://alpha.com/", "state": "Delaware", "city": "Wilmington"},
        {"firm_name": "Alpha Law LLC", "website": "https://www.alpha.com/team", "state": "Delaware", "city": "Dover"},
        {"firm_name": "Alpha & Beta LLP", "website": "https://alphabeta.com/", "state": "Delaware", "city": "Newark"},
        {"firm_name": "Beta Law", "website": "https://alpha.com/other", "state": "Delaware", "city": "Dover"},
    ]
    unique, dupes = deduplicate(recs)
    assert [u["firm_name"] for u in unique] == ["Alpha Law LLP", "Alpha & Beta LLP"]
    assert len(dupes) == 2

def test_store_jsonl_append_and_merge(tmp_path):
    from cornell_ip.store import append_jsonl, load_jsonl, merge_records
    p = str(tmp_path / "raw_records.jsonl")
    append_jsonl(p, {"profile_url": "u1", "firm_name": "Alpha Law LLP"})
    append_jsonl(p, {"profile_url": "u2", "firm_name": ""})
    assert [r["profile_url"] for r in load_jsonl(p)] == ["u1", "u2"]
    legacy = [{"profile_url": "u1", "firm_name": "Old Name"},
              {"profile_url": "u0", "firm_name": "Zero LLP"}]
    merged = merge_records(legacy, load_jsonl(p))
    # union, first-seen order, JSONL version of u1 wins
    assert [r["profile_url"] for r in merged] == ["u1", "u0", "u2"]
    assert merged[0]["firm_name"] == "Alpha Law LLP"


def test_progress_never_overstates_completion():
    from cornell_ip.store import build_progress
    recs = [{"firm_name": "Alpha Law LLP", "website": "https://a.com/",
             "state": "Texas", "city": "Dallas",
             "listing_url": "https://x/texas?page=2", "profile_url": "u1"}]
    sp, cp = build_progress(recs)
    assert len(sp) == 51  # 50 states + District of Columbia
    assert sp["Texas"]["status"] == "in_progress"
    assert sp["Texas"]["listing_pages_processed_count"] == 1
    assert sp["North Dakota"]["status"] == "not_started"
    sp2, _ = build_progress(recs, exhausted_states={"Texas"},
                            no_listing_states={"North Dakota"})
    assert sp2["Texas"]["status"] == "completed"
    assert sp2["North Dakota"]["status"] == "no_relevant_listings_found"
    assert cp["Texas::Dallas"]["firm_bearing_records"] == 1


def test_two_column_workbook_exact_spec(tmp_path):
    pytest.importorskip("openpyxl")
    from openpyxl import load_workbook
    from cornell_ip.excel import build_firm_workbook
    out = str(tmp_path / "output" / "IP_Law_Firms.xlsx")
    firms = [{"firm_name": "Alpha Law LLP", "website": "https://alpha.com/"},
             {"firm_name": "Beta Law", "website": ""}]
    build_firm_workbook(firms, out)
    wb = load_workbook(out)
    assert wb.sheetnames == ["IP Law Firms"]
    ws = wb["IP Law Firms"]
    assert ws.max_column == 2
    assert [ws.cell(1, c).value for c in (1, 2)] == \
        ["Law Firm Name", "Official Website URL"]
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == "A1:B3"
    assert ws.cell(2, 2).hyperlink.target == "https://alpha.com/"
    assert ws.cell(3, 2).value is None  # unverified website left blank


def test_contact_block_firm_and_solo():
    pytest.importorskip("scrapy")
    from scrapy.http import HtmlResponse, Request
    from cornell_ip.spiders.cornell_ip import CornellIpSpider
    spider = CornellIpSpider()
    for fname, expect_firm in [("profile_firm.html", "Morris, Nichols, Arsht & Tunnell LLP"),
                               ("profile_solo.html", "")]:
        body = open(os.path.join(FIX, fname), "rb").read()
        resp = HtmlResponse(url="https://lawyers.law.cornell.edu/lawyer/x-1",
                            body=body, encoding="utf-8",
                            request=Request(url="https://lawyers.law.cornell.edu/lawyer/x-1",
                                            meta={"state": "Delaware", "city_hint": "", "website_hint": "", "listing_url": ""}))
        items = list(spider.parse_profile(resp))
        if expect_firm:
            assert len(items) == 1
            assert items[0]["firm_name"] == expect_firm
            assert items[0]["city"] == "Wilmington" and items[0]["state"] == "Delaware"
            assert items[0]["website"].startswith("https://www.mnat.com/")
        else:
            assert items == []
