"""Build US_IP_Law_Firms_Cornell.xlsx: one sheet 'IP Law Firms', exactly
5 columns, formatted header, filters, frozen header row, clickable URLs."""
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table  # noqa: F401  (kept for compat)

HEADERS = ["Serial Number", "Firm Name", "Official Website URL",
           "U.S. State", "City"]
WIDTHS = [14, 52, 46, 24, 24]


def build_workbook(firms, out_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "IP Law Firms"
    ws.append(HEADERS)

    header_fill = PatternFill("solid", fgColor="1F3864")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin = Side(style="thin", color="B4C6E7")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for c in range(1, 6):
        cell = ws.cell(row=1, column=c)
        cell.fill, cell.font, cell.border = header_fill, header_font, border
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 22

    link_font = Font(name="Calibri", size=11, color="0563C1", underline="single")
    body_font = Font(name="Calibri", size=11)
    alt_fill = PatternFill("solid", fgColor="D9E2F3")
    for i, firm in enumerate(firms, start=1):
        r = i + 1
        ws.cell(row=r, column=1, value=i).font = body_font
        ws.cell(row=r, column=2, value=firm.get("firm_name", "")).font = body_font
        site = (firm.get("website") or "").strip()
        c3 = ws.cell(row=r, column=3, value=site if site else None)
        if site:
            c3.hyperlink = site
            c3.font = link_font
        else:
            c3.font = body_font
        ws.cell(row=r, column=4, value=firm.get("state", "")).font = body_font
        ws.cell(row=r, column=5, value=firm.get("city", "")).font = body_font
        for c in range(1, 6):
            cell = ws.cell(row=r, column=c)
            cell.border = border
            cell.alignment = Alignment(vertical="center",
                                       horizontal="center" if c in (1, 4, 5) else "left")
            if i % 2 == 0:
                cell.fill = alt_fill

    for idx, w in enumerate(WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:E{max(1, len(firms) + 1)}"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.print_title_rows = "1:1"
    wb.save(out_path)
    return out_path


FIRM_HEADERS = ["Law Firm Name", "Official Website URL"]
FIRM_WIDTHS = [56, 62]


def build_firm_workbook(firms, out_path):
    """Final two-column workbook: verified firm name + official website.

    ``firms`` must already be globally deduplicated.  Website cells stay
    blank when no official website could be reliably established; verified
    URLs are clickable hyperlinks.  No other columns or sheets are added.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "IP Law Firms"
    ws.append(FIRM_HEADERS)

    header_fill = PatternFill("solid", fgColor="1F3864")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin = Side(style="thin", color="B4C6E7")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for c in range(1, 3):
        cell = ws.cell(row=1, column=c)
        cell.fill, cell.font, cell.border = header_fill, header_font, border
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 22

    link_font = Font(name="Calibri", size=11, color="0563C1", underline="single")
    body_font = Font(name="Calibri", size=11)
    alt_fill = PatternFill("solid", fgColor="D9E2F3")
    for i, firm in enumerate(firms, start=1):
        r = i + 1
        name_cell = ws.cell(row=r, column=1, value=firm.get("firm_name", ""))
        name_cell.font = body_font
        site = (firm.get("website") or "").strip()
        site_cell = ws.cell(row=r, column=2, value=site if site else None)
        if site:
            site_cell.hyperlink = site
            site_cell.font = link_font
        else:
            site_cell.font = body_font
        for c in range(1, 3):
            cell = ws.cell(row=r, column=c)
            cell.border = border
            cell.alignment = Alignment(vertical="center")
            if i % 2 == 0:
                cell.fill = alt_fill

    for idx, w in enumerate(FIRM_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:B{max(1, len(firms) + 1)}"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.print_title_rows = "1:1"
    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    wb.save(out_path)
    return out_path
